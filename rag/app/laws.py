#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
import copy
import re
from io import BytesIO
from docx import Document
from rag.nlp import bullets_category, is_english, tokenize, remove_contents_table, hierarchical_merge, \
    make_colon_as_title, add_positions
from rag.nlp import huqie
from deepdoc.parser import PdfParser, DocxParser
from rag.settings import cron_logger


class Docx(DocxParser):
    def __init__(self):
        pass

    def __clean(self, line):
        line = re.sub(r"\u3000", " ", line).strip()
        return line

    def __call__(self, filename, binary=None, from_page=0, to_page=100000):
        self.doc = Document(
            filename) if not binary else Document(BytesIO(binary))
        pn = 0
        lines = []
        for p in self.doc.paragraphs:
            if pn > to_page:break
            if from_page <= pn < to_page and p.text.strip(): lines.append(self.__clean(p.text))
            for run in p.runs:
                if 'lastRenderedPageBreak' in run._element.xml:
                    pn += 1
                    continue
                if 'w:br' in run._element.xml and 'type="page"' in run._element.xml:
                    pn += 1
        return [l for l in lines if l]


class Pdf(PdfParser):
    def __call__(self, filename, binary=None, from_page=0,
                 to_page=100000, zoomin=3, callback=None):
        callback(msg="OCR is  running...")
        self.__images__(
            filename if not binary else binary,
            zoomin,
            from_page,
            to_page,
            callback
        )
        callback(msg="OCR finished")

        from timeit import default_timer as timer
        start = timer()
        self._layouts_rec(zoomin)
        callback(0.67, "Layout analysis finished")
        cron_logger.info("paddle layouts:".format((timer()-start)/(self.total_page+0.1)))
        self._naive_vertical_merge()

        callback(0.8, "Text extraction finished")

        return [b["text"] + self._line_tag(b, zoomin) for b in self.boxes]


def chunk(filename, binary=None, from_page=0, to_page=100000, lang="Chinese", callback=None, **kwargs):
    """
        Supported file formats are docx, pdf, txt.
    """
    doc = {
        "docnm_kwd": filename,
        "title_tks": huqie.qie(re.sub(r"\.[a-zA-Z]+$", "", filename))
    }
    doc["title_sm_tks"] = huqie.qieqie(doc["title_tks"])
    pdf_parser = None
    sections = []
    if re.search(r"\.docx?$", filename, re.IGNORECASE):
        callback(0.1, "Start to parse.")
        for txt in Docx()(filename, binary):
            sections.append(txt)
        callback(0.8, "Finish parsing.")
    elif re.search(r"\.pdf$", filename, re.IGNORECASE):
        pdf_parser = Pdf()
        for txt in pdf_parser(filename if not binary else binary,
                         from_page=from_page, to_page=to_page, callback=callback):
            sections.append(txt)
    elif re.search(r"\.txt$", filename, re.IGNORECASE):
        callback(0.1, "Start to parse.")
        txt = ""
        if binary:txt = binary.decode("utf-8")
        else:
            with open(filename, "r") as f:
                while True:
                    l = f.readline()
                    if not l:break
                    txt += l
        sections = txt.split("\n")
        sections = [l for l in sections if l]
        callback(0.8, "Finish parsing.")
    else: raise NotImplementedError("file type not supported yet(docx, pdf, txt supported)")

    # is it English
    eng = lang.lower() == "english"#is_english(sections)
    # Remove 'Contents' part
    remove_contents_table(sections, eng)

    make_colon_as_title(sections)
    bull = bullets_category(sections)
    cks = hierarchical_merge(bull, sections, 3)
    if not cks: callback(0.99, "No chunk parsed out.")

    res = []
    # wrap up to es documents
    for ck in cks:
        print("\n-".join(ck))
        ck = "\n".join(ck)
        d = copy.deepcopy(doc)
        if pdf_parser:
            d["image"], poss = pdf_parser.crop(ck, need_position=True)
            add_positions(d, poss)
            ck = pdf_parser.remove_tag(ck)
        tokenize(d, ck, eng)
        res.append(d)
    return res


if __name__ == "__main__":
    import sys
    def dummy(prog=None, msg=""):
        pass
    chunk(sys.argv[1], callback=dummy)